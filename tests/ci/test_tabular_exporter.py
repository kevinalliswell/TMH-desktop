from __future__ import annotations

import csv

from openpyxl import load_workbook

from src.application.dto import ExperimentDetailDTO
from src.application.services.report_export_service import ReportExportService
from src.services.experiment_file import ExperimentFile
from src.utils.data_saver import DataSaver
from src.utils.tabular_exporter import TabularExporter


def _detail() -> ExperimentDetailDTO:
    return ExperimentDetailDTO(
        experiment_id="EXP-EXPORT",
        experiment_name="中文实验",
        sample_name="样品",
        sample_weight=500.0,
        start_time="2026-08-31T10:00:00",
        end_time="2026-08-31T11:00:00",
        operator="tester",
        experiment_type="GB/T 13241",
        description="export contract",
        timestamps=["2026-08-31T10:00:00"],
        temperatures=[900.0],
        weights=[499.0],
        weight_losses=[0.2],
        gas_flows={"CO": [4.5], "CO2": [], "N2": [10.5], "H2": []},
    )


def test_history_formats_share_one_header_and_missing_value_policy(tmp_path):
    service = ReportExportService(history_query_service=object())
    detail = _detail()
    csv_path = tmp_path / "data.csv"
    txt_path = tmp_path / "data.txt"
    xlsx_path = tmp_path / "data.xlsx"

    service._export_csv(detail, csv_path)
    service._export_txt(detail, txt_path)
    service._export_xlsx(detail, xlsx_path)

    csv_rows = list(csv.reader(csv_path.open(encoding="utf-8-sig")))
    txt_rows = list(csv.reader(txt_path.open(encoding="utf-8"), delimiter="\t"))
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    xlsx_rows = list(workbook["实验数据"].iter_rows(values_only=True))

    expected_header = list(ReportExportService.DATA_HEADERS)
    assert expected_header in csv_rows
    assert expected_header in txt_rows
    assert list(xlsx_rows[0]) == expected_header

    csv_data = csv_rows[csv_rows.index(expected_header) + 1]
    txt_data = txt_rows[txt_rows.index(expected_header) + 1]
    assert csv_data[-3:] == ["", "10.5", ""]
    assert txt_data[-3:] == ["", "10.5", ""]
    assert list(xlsx_rows[1])[-3:] == [None, 10.5, None]
    assert csv_path.read_bytes().startswith(b"\xef\xbb\xbf")


def test_realtime_and_history_exports_use_shared_tabular_writer():
    assert DataSaver.tabular_exporter is TabularExporter
    assert ReportExportService.tabular_exporter is TabularExporter


def test_experiment_file_only_manages_native_experiment_files():
    assert not hasattr(ExperimentFile, "export_data")
    assert not hasattr(ExperimentFile, "generate_report")
